"""Build the fixed core65 and restored74 branches from full KO counts, never relabel counts."""
from pathlib import Path
import csv, gzip, hashlib, json
from collections import Counter, defaultdict
from datetime import datetime, timezone
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'code/result_raw/KO_panel_repair_20260926'
DEST=OUT/'amended_inputs'
EXCLUDE={'K03495','K04566','K00131','K10254','K06719','K08577'}
ADDS={
 'K03573':('dna_repair_radiation','mismatch_repair','Restore MutH identity'),
 'K00836':('osmotic_desiccation_salt','ectoine','Restore ectoine aminotransferase'),
 'K00108':('osmotic_desiccation_salt','betaine','Restore choline oxidation'),
 'K10257':('cold_protein_quality','membrane_adaptation','Restore explicitly named DesB omega-3 role'),
 **{f'K{k}':('dormancy_resuscitation','resuscitation','Include all five official Rpf family members, not an arbitrary one') for k in range(21687,21692)}}
LABELS={
 'dna_repair_radiation':'DNA repair and protection',
 'oxidative_redox':'Oxidative/redox defence',
 'osmotic_desiccation_salt':'Osmotic/solute and ion homeostasis',
 'cold_protein_quality':'Cold-associated RNA and lipid functions',
 'dormancy_resuscitation':'Sporulation and stringent response',
 'trace_gas_energy':'Hydrogen/CO metabolism and carbon fixation',
 'light_energy':'Light-driven ion transport',
 'sulfur_chemolithotrophy':'Sulfur metabolism',
 'biofilm_eps_surface':'Extracellular polysaccharide-associated functions'}
def read(p,sep=','):
 with (gzip.open if str(p).endswith('.gz') else open)(p,'rt',newline='') as f:
  yield from csv.DictReader(f,delimiter=sep)
def write(p,rows,fields=None,sep='\t'):
 with (gzip.open if str(p).endswith('.gz') else open)(p,'wt',newline='') as f:
  w=csv.DictWriter(f,fields or list(rows[0]),delimiter=sep);w.writeheader();w.writerows(rows)
def sha(p): return hashlib.sha256(p.read_bytes()).hexdigest()
def main():
 if DEST.exists(): raise SystemExit('Refusing to overwrite amended_inputs')
 amendment=OUT/'analysis_amendment_before_rerun.md'
 assert amendment.exists()
 audit=list(read(OUT/'historical_71KO_annotation_audit.tsv','\t'))
 old={r['KO']:r for r in audit}
 official={}
 for line in (OUT/'KEGG_official_list_ko_20260926.txt').read_text().splitlines():
  k,v=line.split('\t',1); sym,name=v.split('; ',1) if '; ' in v else ('',v)
  official[k.removeprefix('ko:')]=(sym,name)
 snap={r['knum']:r for r in read(ROOT/'code/result_raw/kegg_module_kofam_observability_20260714_v2/kofam_ko_list_2026-06-01_snapshot.tsv','\t')}
 axes={r['original_module']:r['strategy_axis'] for r in audit}
 branches={}
 for name,kos in [('core65',set(old)-EXCLUDE),('restored74',(set(old)-EXCLUDE)|set(ADDS))]:
  rows=[]
  for ko in sorted(kos):
   if ko in old: m,sub,reason=old[ko]['original_module'],old[ko]['original_submodule'],'Retain KO identity with verified official label'
   else: m,sub,reason=ADDS[ko]
   sym,func=official[ko]
   assert snap[ko]['definition']==func,(ko,'historical/current definition mismatch')
   rows.append(dict(KO=ko,gene=sym,definition=func,strategy_axis=axes[m],module=m,
     module_display=LABELS[m]+(' and Rpf' if m=='dormancy_resuscitation' and name=='restored74' else ''),
     submodule=sub,decision=reason,source=f'https://www.kegg.jp/entry/{ko}'))
  assert len(rows)==int(name[-2:]); branches[name]=rows
 selected=set(r['KO'] for rs in branches.values() for r in rs)
 counts={}
 for r in read(ROOT/'code/result_raw/kofam_primary_aggregation_parallel_20260714/primary_3904_mag_ko_counts.tsv.gz','\t'):
  if r['KO'] in selected:
   key=(r['dataset_id'],r['genome_id'],r['KO']);assert key not in counts
   counts[key]=int(r['protein_count'])
 original=ROOT/'code/result_raw/whole_kofam_prespecified_71KO_20260714/primary_3904_prespecified_71KO_presence.csv.gz'
 meta={}; mismatches=0
 fields=None
 for r in read(original):
  fields=fields or list(r)
  key=(r['dataset_id'],r['genome_id'])
  meta.setdefault(key,{k:v for k,v in r.items() if k not in ['strategy_axis','module','submodule','KO','gene','protein_count','present']})
  if r['KO'] in selected: mismatches+=int(int(r['protein_count'])!=counts.get((*key,r['KO']),0))
 assert len(meta)==3904 and mismatches==0
 DEST.mkdir()
 stamp={'created_utc':datetime.now(timezone.utc).isoformat(),'amendment_sha256':sha(amendment),
        'amendment':'post-audit correction, not retrospective prespecification',
        'unchanged_KO_count_mismatches':mismatches,'branches':{}}
 wb=openpyxl.Workbook(); ws=wb.active;ws.title='Read_me'
 for row in [
  ['Supplementary Table S9 — annotation-audited mapping; downstream reruns pending'],
  ['core65','Minimal corrected panel: 65 KOs, nine modules, three broad strategy groups.'],
  ['restored74','Functional-restoration sensitivity: 74 KOs, nine modules, three broad strategy groups.'],
  ['Scope','Names and functions are official KO definitions, not measurements of activity or complete pathways.'],
  ['Amendment','Corrects the historical 71-KO configuration after annotation audit. The corrected panels were not publicly preregistered.'],
  ['Breadth','Fraction of branch-specific panel KOs detected within each module.'],
  ['Rpf','All five official Rpf members are included in restored74; this gives the family a larger within-module weight.'],
  ['Energy-associated','Hydrogen metabolism/carbon fixation, sulfur metabolism and ion pumping do not establish atmospheric affinity, net energy harvesting or lithotrophy.'],
  ['Status','Mapping is verified. Do not pair these tables with unrevised historical figure statistics.']]:ws.append(row)
 for name,rows in branches.items():
  key={r['KO']:r for r in rows}
  panelpath=DEST/(name+'_panel.tsv');write(panelpath,rows)
  presencepath=DEST/(name+'_presence.csv.gz')
  with gzip.open(presencepath,'wt',newline='') as f:
   w=csv.DictWriter(f,fields);w.writeheader()
   for mk in sorted(meta):
    for ko,p in key.items():
     c=counts.get((*mk,ko),0)
     w.writerow(dict(meta[mk],strategy_axis=p['strategy_axis'],module=p['module'],submodule=p['submodule'],KO=ko,gene=p['gene'],protein_count=c,present=c>0))
  # One row per MAG/module: explicit denominators and exact original covariates.
  mods=defaultdict(list)
  for r in rows:mods[r['module']].append(r['KO'])
  breadth=[]
  for mk in sorted(meta):
   for m,kos in mods.items():
    n=sum(counts.get((*mk,k),0)>0 for k in kos)
    breadth.append(dict(meta[mk],module=m,strategy_axis=axes[m],n_panel_KOs=len(kos),n_detected_KOs=n,module_breadth=n/len(kos)))
  write(DEST/(name+'_MAG_module_breadth.tsv.gz'),breadth)
  s=wb.create_sheet(name+'_KO_mapping');s.append(list(rows[0]))
  for r in rows:s.append(list(r.values()))
  ss=wb.create_sheet(name+'_module_summary');ss.append(['module','strategy_axis','panel_KOs'])
  for m,ks in mods.items(): ss.append([m,axes[m],len(ks)])
  stamp['branches'][name]={'KOs':len(rows),'modules':len(mods),'strategy_counts':dict(Counter(r['strategy_axis'] for r in rows)),
   'panel_sha256':sha(panelpath),'presence_sha256':sha(presencepath),'rows':3904*len(rows)}
 for sheet in wb:
  sheet.freeze_panes='A2';sheet.auto_filter.ref=sheet.dimensions
  for cell in sheet[1]:cell.font=Font(name='Arial',bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='29465B')
  for row in sheet.iter_rows(min_row=2):
   for cell in row:cell.font=Font(name='Arial',size=10);cell.alignment=Alignment(wrap_text=True,vertical='top')
  for col in sheet.columns:sheet.column_dimensions[col[0].column_letter].width=min(70,max(16,max(len(str(c.value or '')) for c in col)+2))
 wb.save(DEST/'Table_S9_annotation_corrected_mapping_reruns_pending.xlsx')
 (DEST/'build_validation.json').write_text(json.dumps(stamp,indent=2))
 print(json.dumps(stamp,indent=2))
if __name__=='__main__':main()

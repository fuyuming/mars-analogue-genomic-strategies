"""Audit the historical KO panel; stage corrections without altering original analyses.

Run from any directory. Downloads the official KEGG KO list once, compares it
with the June KOfam snapshot and student S9, verifies original presence counts,
and exports a clearly labelled six-exclusion diagnostic branch. No model fitting.
"""
from pathlib import Path
import csv
import gzip
import hashlib
import json
import re
import urllib.request
from collections import Counter, defaultdict
from datetime import datetime, timezone
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
import zipfile

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'code/result_raw/KO_panel_repair_20260926'
OUT.mkdir(parents=True, exist_ok=True)
PANEL = ROOT / 'code/dataset/mars_analog_target_ko_modules.csv'
SNAP = ROOT / 'code/result_raw/kegg_module_kofam_observability_20260714_v2/kofam_ko_list_2026-06-01_snapshot.tsv'
AGG = ROOT / 'code/result_raw/kofam_primary_aggregation_parallel_20260714'
PRES = ROOT / 'code/result_raw/whole_kofam_prespecified_71KO_20260714/primary_3904_prespecified_71KO_presence.csv.gz'
STUDENT = ROOT / 'srx-meta-ysc-9.22'
AXES = {m: a for a, ms in {
    'cell_maintenance': ['dna_repair_radiation', 'oxidative_redox', 'cold_protein_quality', 'osmotic_desiccation_salt', 'dormancy_resuscitation'],
    'energy_acquisition': ['trace_gas_energy', 'light_energy', 'sulfur_chemolithotrophy'],
    'surface_retention': ['biofilm_eps_surface']}.items() for m in ms}
EXCLUDE = {
    'K03495': 'tRNA modification does not implement MutH mismatch repair',
    'K04566': 'class-I lysyl-tRNA synthetase is not superoxide dismutase',
    'K00131': 'GapN is not betaine-aldehyde dehydrogenase',
    'K10254': 'oleate hydratase is not a fatty-acid desaturase',
    'K06719': 'CD300 antigen is not ectoine-biosynthesis aminotransferase',
    'K08577': 'calpain-8 is not a bacterial resuscitation-promoting factor',
}
CONTEXT = {
    'trace_gas_energy': 'Hydrogen/CO-associated metabolism and carbon fixation; neither atmospheric affinity nor net energy acquisition follows from these KOs alone.',
    'sulfur_chemolithotrophy': 'Sulfur-metabolism-associated markers; DsrAB alone does not determine oxidation versus reduction or lithotrophy.',
    'light_energy': 'Light-driven ion transport; chloride-pumping halorhodopsin is not equivalent to proton-driven ATP generation.',
    'dormancy_resuscitation': 'Sporulation/stringent-response markers remain after exclusion; resuscitation is not evaluated by this retained subset.',
}

def read(path, sep=','):
    op = gzip.open if str(path).endswith('.gz') else open
    with op(path, 'rt', newline='', encoding='utf-8-sig') as f:
        return list(csv.DictReader(f, delimiter=sep))

def write(name, rows, fields=None):
    fields = fields or list(rows[0])
    path = OUT / name
    op = gzip.open if name.endswith('.gz') else open
    with op(path, 'wt', newline='', encoding='utf-8') as f:
        w = csv.DictWriter(f, fields, delimiter='\t' if '.tsv' in name else ',')
        w.writeheader(); w.writerows(rows)

def digest(path):
    h = hashlib.sha256()
    with open(path, 'rb') as f:
        for b in iter(lambda: f.read(1048576), b''): h.update(b)
    return h.hexdigest()

def main():
    panel = read(PANEL)
    assert len(panel) == len({r['ko'] for r in panel}) == 71
    snap = {r['knum']: r for r in read(SNAP, '\t')}
    cache = OUT / 'KEGG_official_list_ko_20260926.txt'
    if not cache.exists():
        req = urllib.request.Request('https://rest.kegg.jp/list/ko', headers={'User-Agent': 'ISME-KO-annotation-audit/1.0'})
        with urllib.request.urlopen(req, timeout=60) as response:
            raw = response.read()
        assert raw.startswith(b'K00001\t') or raw.startswith(b'ko:K00001\t')
        cache.write_bytes(raw)
        (OUT / 'official_retrieval.json').write_text(json.dumps({
            'url': req.full_url, 'retrieved_utc': datetime.now(timezone.utc).isoformat(),
            'sha256': digest(cache)}, indent=2))
    official = {}
    for line in cache.read_text().splitlines():
        ko, desc = line.split('\t', 1); ko = ko.removeprefix('ko:')
        symbols, name = desc.split('; ', 1) if '; ' in desc else ('', desc)
        official[ko] = {'symbols': symbols, 'name': name}
    s9 = STUDENT / 'Supplementary_Tables/Supplementary_Table_S9.xlsx'
    wb = openpyxl.load_workbook(s9, read_only=True, data_only=True)
    sheets = {}
    for ws in wb:
        ws.reset_dimensions(); sheets[ws.title] = list(ws.values)
    smap = {r[2]: r for r in sheets['KO_mapping'][1:] if len(r) >= 8 and r[2]}
    flagged = {r[0] for r in sheets['Annotation_audit'][1:] if r and r[0]}
    assert len(smap) == 71 and len(flagged) == 22
    audit = []
    for p in panel:
        ko = p['ko']; ref = official[ko]; s = smap[ko]
        audit.append({
            'KO': ko, 'original_gene': p['gene'], 'official_symbols': ref['symbols'],
            'official_name': ref['name'], 'original_module': p['module'],
            'original_submodule': p['submodule'], 'strategy_axis': AXES[p['module']],
            'original_description': p['description'], 'June_KOfam_definition': snap[ko]['definition'],
            'student_name_matches_official': s[3] == ref['name'],
            'student_symbols_match_official': s[6] == ref['symbols'],
            'June_definition_matches_current': snap[ko]['definition'] == ref['name'],
            'student_flagged': ko in flagged,
            'action': 'EXCLUDE_FROM_SIX_EXCLUSION_DIAGNOSTIC' if ko in EXCLUDE else 'CORRECT_LABEL_AND_INTERPRETATION' if ko in flagged else 'KEEP_KO_WITH_OFFICIAL_LABEL',
            'reason_or_boundary': EXCLUDE.get(ko, CONTEXT.get(p['module'], 'KO-level potential, not pathway completeness or demonstrated activity.')),
            'KEGG_entry': f'https://www.kegg.jp/entry/{ko}',
        })
    write('historical_71KO_annotation_audit.tsv', audit)
    # Recovery probes are NOT additions to the corrected panel. Rpf is a family,
    # not a unique replacement identifier; report all five official Rpf entries.
    probes = {'K03573': 'MutH', 'K00836': 'EctB', 'K00108': 'choline dehydrogenase candidate',
              **{f'K2168{i}': f'Rpf{chr(65+i-7)}' for i in [7,8,9]},
              'K21690': 'RpfD', 'K21691': 'RpfE'}
    selected = {r['ko'] for r in panel} | set(probes)
    subset = [r for r in read(AGG / 'primary_3904_mag_ko_counts.tsv.gz', '\t') if r['KO'] in selected]
    counts = {(r['dataset_id'], r['genome_id'], r['KO']): int(r['protein_count']) for r in subset}
    presence = read(PRES)
    assert len(presence) == 3904 * 71
    mismatches = [r for r in presence if int(r['protein_count']) != counts.get((r['dataset_id'], r['genome_id'], r['KO']), 0)]
    assert not mismatches, f'{len(mismatches)} original presence cells differ from KO aggregation'
    corr = {r['KO']: r for r in audit}
    diagnostic = []
    for r in presence:
        if r['KO'] in EXCLUDE: continue
        s = dict(r); s['gene'] = official[r['KO']]['symbols']; diagnostic.append(s)
    assert len(diagnostic) == 3904 * 65
    write('diagnostic65_presence.csv.gz', diagnostic)
    write('diagnostic65_panel.csv', [{**p, 'gene': official[p['ko']]['symbols'], 'description': official[p['ko']]['name']} for p in panel if p['ko'] not in EXCLUDE])
    write('recovered_candidate_MAG_KO_counts.tsv.gz', [r for r in subset if r['KO'] in probes], list(subset[0]))
    observed = defaultdict(list)
    for r in subset: observed[r['KO']].append(r)
    evidence = []
    for ko in sorted(selected):
        rr = observed[ko]
        evidence.append({'KO': ko, 'official_symbols': official[ko]['symbols'], 'official_name': official[ko]['name'],
            'is_original_panel': ko not in probes, 'recovery_probe_only': ko in probes,
            'primary_MAGs_with_KO': len(rr), 'ANI95_representatives_with_KO': sum(r['species_representative'].lower() == 'true' for r in rr),
            'protein_KO_assignments': sum(int(r['protein_count']) for r in rr),
            'sources_with_KO': len({r['dataset_id'] for r in rr}),
            'interpretation': 'Candidate retrieval only; not a final replacement panel' if ko in probes else corr[ko]['action']})
    write('KO_evidence_and_recovery_probes.tsv', evidence)
    # Repair historical WGS denominator reporting, retaining original measurements.
    wgs = read(STUDENT / 'Source_Data/independent_wgs/KO_detection.tsv', '\t')
    completion = []
    gates = sorted({r['gate'] for r in wgs})
    pby = {p['ko']: p for p in panel}
    absent = set(pby) - {r['KO'] for r in wgs}
    assert absent == {'K06719', 'K08577'}
    for gate in gates:
        for ko in sorted(pby):
            found = [r for r in wgs if r['gate'] == gate and r['KO'] == ko]
            assert len(found) <= 1
            p = pby[ko]
            if found:
                r = dict(found[0]); r['gene'] = official[ko]['symbols']
                r['reference_status'] = 'represented_in_historical_reference'
            else:
                r = {k: 'NA' for k in wgs[0]}
                r.update(gate=gate, KO=ko, gene=official[ko]['symbols'], strategy_axis=AXES[p['module']], module=p['module'], submodule=p['submodule'])
                r['reference_status'] = 'not_represented_not_evaluable'
            r['original_gene_label'] = p['gene']
            r['panel_correction_status'] = corr[ko]['action']
            completion.append(r)
    write('S7_historical71KO_detection_complete.tsv', completion)
    summaries = []
    for gate in gates:
        rr = [r for r in wgs if r['gate'] == gate]
        dd = [r for r in rr if r['KO'] not in EXCLUDE]
        for label, group, size in [('historical71', rr, 71), ('six_exclusion_diagnostic65', dd, 65)]:
            summaries.append({'panel': label, 'gate': gate, 'panel_KOs': size, 'reference_represented_KOs': len(group),
                'not_evaluable_KOs': size-len(group),
                'KOs_detected_in_any_library_both_mates': sum(int(r['samples_mate_replicated']) > 0 for r in group)})
    write('WGS_denominator_audit.tsv', summaries)
    # Stable tabular handoff: no hidden sheets, stale drawings or formulas. This
    # is an audit supplement, explicitly not a publication-ready revised S9.
    outwb = openpyxl.Workbook(); outwb.remove(outwb.active)
    tables = {'README': [
        {'item': 'Status', 'value': 'ANNOTATION REPAIR AUDIT; not a final submission S9'},
        {'item': 'Original panel', 'value': '71 identifiers retained in historical audit; 22 student-flagged discrepancies; six inappropriate module assignments excluded only in diagnostic branch.'},
        {'item': 'Counts', 'value': 'All 277184 original MAG x KO protein counts checked against the full-KOfam KO aggregation.'},
        {'item': 'S7 NA', 'value': 'K06719 and K08577 had no represented reference proteins. WGS evidence is not evaluable, not a measured biological zero.'},
        {'item': 'Inference', 'value': 'Corrected-label raw data do not validate the original module results. All affected fits and BH families require an explicitly amended analysis.'},
        {'item': 'Prespecification', 'value': 'The historical panel was frozen before its test; these corrections are a documented post-audit amendment, not public preregistration or retroactive prespecification.'},
        {'item': 'Module breadth', 'value': 'Fraction of panel KOs detected within each module; not pathway completeness or activity.'}],
        'Historical71_audit': audit, 'Recovery_probes': [r for r in evidence if r['recovery_probe_only']],
        'S7_all71_status': completion, 'WGS_denominators': summaries,
        'Diagnostic65_mapping': [r for r in audit if r['KO'] not in EXCLUDE]}
    for name, rows in tables.items():
        ws = outwb.create_sheet(name); ws.append(list(rows[0]))
        for r in rows: ws.append(list(r.values()))
        ws.freeze_panes = 'A2'; ws.auto_filter.ref = ws.dimensions
        for cell in ws[1]: cell.font = Font(name='Arial', bold=True, color='FFFFFF'); cell.fill = PatternFill('solid', fgColor='29465B')
        for row in ws.iter_rows(min_row=2):
            for cell in row: cell.font = Font(name='Arial', size=10); cell.alignment = Alignment(vertical='top', wrap_text=True)
        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width = min(60, max(16, max(len(str(c.value or '')) for c in col[:40])+2))
    outwb.save(OUT / 'KO_panel_repair_audit_NOT_FINAL_S9.xlsx')
    integrity = {'original_presence_cells_verified': len(presence), 'original_presence_count_mismatches': len(mismatches),
        'student_discrepancy_rows': len(flagged), 'student_function_mismatches_current': [r['KO'] for r in audit if not r['student_name_matches_official']],
        'student_symbols_mismatches_current': [r['KO'] for r in audit if not r['student_symbols_match_official']],
        'June_vs_current_definition_changes': [r['KO'] for r in audit if not r['June_definition_matches_current']],
        'original_unrepresented_KOs': sorted(absent), 'six_exclusion_diagnostic_rows': len(diagnostic),
        'inputs_sha256': {str(p.relative_to(ROOT)): digest(p) for p in [PANEL, SNAP, PRES, s9, AGG/'primary_3904_mag_ko_counts.tsv.gz']}}
    with zipfile.ZipFile(s9) as z:
        xmls = '\n'.join(z.read(n).decode('utf-8', errors='replace') for n in z.namelist() if n.endswith('.xml'))
        integrity['S9_AI_or_placeholder_terms'] = sorted(set(re.findall(r'ChatGPT|OpenAI|Codex|AI-generated|TODO|TO BE PROVIDED', xmls, flags=re.I)))
        integrity['S9_comments_files'] = [n for n in z.namelist() if 'comment' in n.lower()]
    integrity['S9_hidden_sheets'] = [s.title for s in wb if s.sheet_state != 'visible']
    (OUT/'audit_integrity.json').write_text(json.dumps(integrity, indent=2))
    print(json.dumps(integrity, indent=2)); print('WGS', json.dumps(summaries))

if __name__ == '__main__': main()

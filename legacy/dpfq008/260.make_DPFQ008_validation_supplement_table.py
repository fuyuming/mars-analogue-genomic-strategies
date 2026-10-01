#!/usr/bin/env python3
"""Build non-overwriting Supplementary Table S6 for DPFQ008 validation."""

from __future__ import annotations

import csv
import hashlib
import shutil
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo


ROOT = Path(__file__).resolve().parents[1]
OUTDIR = ROOT / "code/result_raw/ISME_Result_DPFQ008_validation_20260723_v1"
OUTPUT = OUTDIR / "Supplementary_Table_S6_DPFQ008_validation.xlsx"
STAGE = ROOT / "ISME_submission_release_workspace_20260718_v1/supplementary/tables/tableS6_DPFQ008_validation.xlsx"

INPUTS = {
    "architecture": ROOT / "code/result_raw/ISME_Figure3_DPFQ008_lineage_context_20260723_v8/Source_Data/Figure3a_strict_hit_architecture.csv",
    "tip_tracks": ROOT / "code/result_raw/ISME_Figure3_DPFQ008_lineage_context_20260723_v8/Source_Data/Figure3b_gene_tree_tip_tracks.csv",
    "fold_summary": ROOT / "code/result_raw/DPFQ008_hmm_seed_cv_20260723_v1/fold_summary.tsv",
    "fold_frequency": ROOT / "code/result_raw/DPFQ008_hmm_seed_cv_20260723_v1/hit_fold_frequency.tsv",
    "gene_host": ROOT / "code/result_raw/DPFQ008_gene_host_tree_congruence_20260723_v3/DPFQ008_gene_host_tree_congruence.tsv",
    "tree_sensitivity": ROOT / "code/result_raw/DPFQ008_gene_host_tree_congruence_20260723_v3/DPFQ008_Qaidam_dominance_sensitivity.tsv",
    "tree_tips": ROOT / "code/result_raw/DPFQ008_gene_host_tree_congruence_20260723_v3/DPFQ008_tree_tip_metadata.tsv",
    "nearest_neighbors": ROOT / "code/result_raw/DPFQ008_gene_host_tree_congruence_20260723_v3/DPFQ008_nearest_neighbor_audit.tsv",
    "matched_test": ROOT / "code/result_raw/DPFQ008_strict63_matched_context_20260722_v6/matched_permutation_summary.tsv",
    "matched_strata": ROOT / "code/result_raw/DPFQ008_strict63_matched_context_20260722_v6/matched_target_control_strata.tsv",
    "neighborhoods": ROOT / "code/result_raw/DPFQ008_strict63_matched_context_20260722_v6/expanded_target_neighborhoods.tsv",
}


def read_table(path: Path) -> list[dict[str, str]]:
    delimiter = "\t" if path.suffix == ".tsv" else ","
    with path.open(newline="", encoding="utf-8") as handle:
        return list(csv.DictReader(handle, delimiter=delimiter))


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def add_sheet(wb: Workbook, name: str, rows: list[dict[str, object]]) -> None:
    if not rows:
        raise ValueError(f"No rows for sheet {name}")
    ws = wb.create_sheet(name)
    headers = list(rows[0])
    ws.append(headers)
    for row in rows:
        if list(row) != headers:
            raise ValueError(f"Inconsistent columns in sheet {name}")
        ws.append([row[key] for key in headers])

    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in ws[1]:
        cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name="Arial", size=9)
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    table = Table(displayName=f"S6{name.replace('_', '')}", ref=ws.dimensions)
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    ws.add_table(table)
    ws.row_dimensions[1].height = 30
    for index, header in enumerate(headers, start=1):
        values = [str(header)] + ["" if row[header] is None else str(row[header]) for row in rows]
        width = min(50, max(10, max(len(value) for value in values) + 2))
        ws.column_dimensions[get_column_letter(index)].width = width


def main() -> None:
    missing = [str(path) for path in INPUTS.values() if not path.exists()]
    if missing:
        raise FileNotFoundError("Missing inputs:\n" + "\n".join(missing))
    if OUTPUT.exists() or STAGE.exists():
        raise FileExistsError(f"Refusing to overwrite: {OUTPUT if OUTPUT.exists() else STAGE}")

    tables = {name: read_table(path) for name, path in INPUTS.items()}
    architecture = {row["protein_id"]: row for row in tables["architecture"]}
    tracks = {row["protein_id"]: row for row in tables["tip_tracks"]}
    if len(architecture) != 63 or set(architecture) != set(tracks):
        raise AssertionError("Strict architecture and tree tracks must contain the same 63 proteins")

    strict_hits = []
    for protein_id in sorted(architecture):
        arch = architecture[protein_id]
        track = tracks[protein_id]
        strict_hits.append(
            {
                "protein_id": protein_id,
                "dataset_id": track["dataset_id"],
                "source": track["source"],
                "genome_id": track["genome_id"],
                "ani_cluster_95": track["ani_cluster_95"],
                "gtdb_taxonomy_r226": track["gtdb_taxonomy_r226"],
                "length_aa": arch["length_aa"],
                "cxxch_count": arch["cxxch_count"],
                "cxxch_positions": arch["cxxch_positions"],
                "domain_ievalue": arch["domain_ievalue"],
                "target_coverage": arch["target_coverage"],
                "target_context_positive": track["target_context_positive"],
            }
        )

    if len({row["genome_id"] for row in strict_hits}) != 61:
        raise AssertionError("Expected 61 carrier MAGs")
    if len({row["ani_cluster_95"] for row in strict_hits}) != 58:
        raise AssertionError("Expected 58 ANI95 clusters")
    source_counts: dict[str, int] = {}
    for row in strict_hits:
        source_counts[row["dataset_id"]] = source_counts.get(row["dataset_id"], 0) + 1
    expected_sources = {
        "qaidam_basin_mmag_1773": 56,
        "mauna_loa_lava_tube_fishman_2023": 3,
        "australian_basalt_lava_tubes_bay_2025": 2,
        "alaska_permafrost_reference": 2,
    }
    if source_counts != expected_sources:
        raise AssertionError(f"Unexpected source counts: {source_counts}")

    if len(tables["fold_summary"]) != 5:
        raise AssertionError("Expected five HMM sensitivity folds")
    if sum(int(row["heldout_recovered"]) for row in tables["fold_summary"]) != 20:
        raise AssertionError("Expected recovery of all 20 held-out seeds")
    full_frequency = [row for row in tables["fold_frequency"] if row["in_full_strict63"] == "1"]
    if len(full_frequency) != 63:
        raise AssertionError("Expected 63 full-profile hits in fold-frequency table")
    if sum(int(row["folds_detected"]) >= 4 for row in full_frequency) != 60:
        raise AssertionError("Expected 60 of 63 hits in at least four folds")

    matched = {row["level"]: row for row in tables["matched_test"]}
    if matched["MAG"]["observed_positive"] != "53" or matched["MAG"]["denominator"] != "61":
        raise AssertionError("Expected 53/61 matched-context MAG result")
    if matched["protein"]["observed_positive"] != "54" or matched["protein"]["denominator"] != "63":
        raise AssertionError("Expected 54/63 matched-context protein result")
    if len(tables["tree_tips"]) != 58 or len(tables["nearest_neighbors"]) != 58:
        raise AssertionError("Expected 58 ANI95 representatives in tree sheets")

    readme = [
        {"item": "Table", "description": "Supplementary Table S6 | Catalogue-wide DPFQ008 validation"},
        {"item": "Purpose", "description": "Profile-HMM, seed-stability, phylogenetic and within-MAG genomic-context evidence for the AI-prioritized DPFQ008 family."},
        {"item": "Strict family", "description": "63 proteins from 61 carrier MAGs and 58 ANI95 clusters across four source catalogues."},
        {"item": "Interpretation", "description": "Lineage-structured electron-transfer-associated candidate family with non-random cytochrome/haem context."},
        {"item": "Boundary", "description": "Does not establish substrate, donor–acceptor pairing, expression, flux, biochemical activity, a new pathway, adaptive convergence or exclusively vertical/horizontal inheritance."},
        {"item": "Discovery table", "description": "Supplementary Table S4 retains the frozen 13-candidate discovery audit and is not replaced by this validation table."},
    ]
    provenance = []
    for name, path in INPUTS.items():
        rows = tables[name]
        provenance.append(
            {
                "input_name": name,
                "relative_path": str(path.relative_to(ROOT)),
                "sha256": sha256(path),
                "bytes": path.stat().st_size,
                "parsed_rows": len(rows),
                "columns": ";".join(rows[0]) if rows else "",
            }
        )

    wb = Workbook()
    wb.remove(wb.active)
    add_sheet(wb, "README", readme)
    add_sheet(wb, "Strict_hits", strict_hits)
    add_sheet(wb, "Seed_CV", tables["fold_summary"])
    add_sheet(wb, "Hit_fold_frequency", tables["fold_frequency"])
    add_sheet(wb, "Gene_host_test", tables["gene_host"])
    add_sheet(wb, "Tree_sensitivity", tables["tree_sensitivity"])
    add_sheet(wb, "Tree_tips", tables["tree_tips"])
    add_sheet(wb, "Nearest_neighbors", tables["nearest_neighbors"])
    add_sheet(wb, "Matched_test", tables["matched_test"])
    add_sheet(wb, "Matched_strata", tables["matched_strata"])
    add_sheet(wb, "Neighborhoods", tables["neighborhoods"])
    add_sheet(wb, "Provenance", provenance)
    wb.properties.title = "Supplementary Table S6 | DPFQ008 validation"
    wb.properties.subject = "ISME Journal internal author draft"
    wb.properties.creator = "Author team"

    OUTDIR.mkdir(parents=True, exist_ok=True)
    STAGE.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUTPUT)

    check = load_workbook(OUTPUT, data_only=False, read_only=True)
    formulas = []
    errors = []
    for ws in check.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                value = cell.value
                if isinstance(value, str) and value.startswith("="):
                    formulas.append(f"{ws.title}!{cell.coordinate}")
                if isinstance(value, str) and value in {"#REF!", "#DIV/0!", "#VALUE!", "#N/A", "#NAME?"}:
                    errors.append(f"{ws.title}!{cell.coordinate}:{value}")
    if formulas or errors:
        raise RuntimeError(f"Unexpected formulas/errors: formulas={formulas}, errors={errors}")
    shutil.copy2(OUTPUT, STAGE)
    print(OUTPUT)
    print(STAGE)


if __name__ == "__main__":
    main()

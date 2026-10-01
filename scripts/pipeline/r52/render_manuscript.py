from pathlib import Path
import re,json
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.section import WD_SECTION_START,WD_ORIENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from PIL import Image
H=Path(__file__).resolve().parent;D=H/'manuscript';s=(D/'manuscript_v6.md').read_text();doc=Document();doc.core_properties.title=s.splitlines()[0][2:];doc.core_properties.author=''
sec=doc.sections[0];sec.page_width=Inches(8.27);sec.page_height=Inches(11.69)
for a in ['top_margin','bottom_margin','left_margin','right_margin']:setattr(sec,a,Inches(.78))
styles=doc.styles
for name in ['Normal','Title','Heading 1','Heading 2','Caption']:
 styles[name].font.name='Arial';styles[name].font.color.rgb=RGBColor(0,0,0)
styles['Normal'].font.size=Pt(10.5);styles['Normal'].paragraph_format.line_spacing=1.12;styles['Normal'].paragraph_format.space_after=Pt(6)
styles['Title'].font.size=Pt(21);styles['Title'].paragraph_format.space_after=Pt(12)
for name,size in [('Heading 1',14),('Heading 2',11.5)]:
 styles[name].font.size=Pt(size);styles[name].paragraph_format.space_before=Pt(11);styles[name].paragraph_format.space_after=Pt(6);styles[name].paragraph_format.keep_with_next=True
styles['Caption'].font.size=Pt(10);styles['Caption'].font.italic=False;styles['Caption'].font.bold=False
for style in styles:
 for el in list(style.element.iter(qn('w:pBdr'))):el.getparent().remove(el)
bib=styles.add_style('ReferencesText',WD_STYLE_TYPE.PARAGRAPH);bib.base_style=styles['Normal'];bib.font.size=Pt(10);bib.paragraph_format.line_spacing=1.05;bib.paragraph_format.space_after=Pt(4)
# Numeric bibliography navigation; no fabricated reference manager metadata.
def bookmark(p,name,i):
 b=OxmlElement('w:bookmarkStart');b.set(qn('w:id'),str(i));b.set(qn('w:name'),name);p._p.append(b)
 e=OxmlElement('w:bookmarkEnd');e.set(qn('w:id'),str(i));p._p.append(e)
def link(p,text,anchor=None,url=None):
 h=OxmlElement('w:hyperlink')
 if anchor:h.set(qn('w:anchor'),anchor)
 if url:h.set(qn('r:id'),p.part.relate_to(url,RT.HYPERLINK,is_external=True))
 r=OxmlElement('w:r');pr=OxmlElement('w:rPr');color=OxmlElement('w:color');color.set(qn('w:val'),'163D55');pr.append(color);r.append(pr);t=OxmlElement('w:t');t.text=text;r.append(t);h.append(r);p._p.append(h)
def text(p,txt):
 for part in re.split(r'(\[\d+(?:,\s*\d+)*\]|https?://\S+)',txt):
  if re.fullmatch(r'\[\d+(?:,\s*\d+)*\]',part):
   p.add_run('[')
   for j,n in enumerate(re.findall(r'\d+',part)):
    if j:p.add_run(', ')
    link(p,n,anchor='ref'+n)
   p.add_run(']')
  elif part.startswith('http'):link(p,part,url=part.rstrip('.'))
  else:p.add_run(part)
# Page numbering is inherited by later sections.
f=sec.footer.paragraphs[0];f.alignment=WD_ALIGN_PARAGRAPH.RIGHT;r=f.add_run();fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');r._r.addnext(fld)
blocks=[b.strip() for b in s.split('\n\n') if b.strip()];figs=[];infig=False;refs=False
for b in blocks:
 if b=='## Figures and legends':infig=True;continue
 if b=='## References':infig=False;refs=True
 if infig:figs.append(b);continue
 if b=='## Acknowledgements':
  for note in blocks:
   if note.startswith('Supplementary Table'):
    p=doc.add_paragraph();text(p,note)
 if b.startswith('# '):doc.add_paragraph(b[2:],style='Title');continue
 if b.startswith('## '):doc.add_paragraph(b[3:],style='Heading 1');continue
 if b.startswith('### '):doc.add_paragraph(b[4:],style='Heading 2');continue
 p=doc.add_paragraph();text(p,b.replace('\n',' '))
 if b.startswith('Descriptive degrees-of-freedom'):p.paragraph_format.keep_together=True
 if refs:
  p.style='ReferencesText'
  m=re.match(r'(\d+)\.',b)
  if m:bookmark(p,'ref'+m[1],100+int(m[1]))
  p.paragraph_format.space_after=Pt(4)
figs=[b for b in figs if not b.startswith("Supplementary Table")]
# Full-width landscape figures preserve multi-panel typography.
figure_count=0
for b in figs:
 if b.startswith('!['):
  sec=doc.add_section(WD_SECTION_START.NEW_PAGE);sec.orientation=WD_ORIENT.LANDSCAPE;sec.page_width=Inches(11.69);sec.page_height=Inches(8.27)
  sec.top_margin=sec.bottom_margin=Inches(.65);sec.left_margin=sec.right_margin=Inches(.75)
  m=re.match(r'!\[(.*?)\]\((.*?)\)',b);im=D/m[2];w,h=Image.open(im).size;width=min(10.0,5.35*w/h)
  if 'figure6_' in str(im):
   sec.page_height=Inches(10.5);width=10.0
  p=doc.add_paragraph();p.paragraph_format.keep_with_next=True;p.paragraph_format.space_after=Pt(9);p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.add_run().add_picture(str(im),width=Inches(width));figure_count+=1
 elif b.startswith('Figure ') or b.startswith('Supplementary Figure'):
  p=doc.add_paragraph(style='Caption');text(p,b)
 else:
  p=doc.add_paragraph();text(p,b)
output=D/'manuscript_v6.docx';doc.save(output)
(D/'build_summary.json').write_text(json.dumps({'source':'manuscript_v6.md','figures':figure_count,'bibliography_entries':len(re.findall(r'^\d+\. ',s,re.M)),'native_internal_citation_links':True,'layout':'editorial reading draft; not claimed official journal template'},indent=2))
print(output)
